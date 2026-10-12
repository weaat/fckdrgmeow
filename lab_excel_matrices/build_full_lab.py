import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()
wb.remove(wb.active) # Удаляем дефолтный пустой лист

# Цвета и шрифты
font_title = Font(name="Calibri", size=14, bold=True, color="1E3A8A")
font_subtitle = Font(name="Calibri", size=11, bold=True, color="1F2937")
font_section = Font(name="Calibri", size=11, bold=True, color="111827")
font_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
font_bold = Font(name="Calibri", size=10, bold=True)
font_regular = Font(name="Calibri", size=10)
font_formula = Font(name="Consolas", size=9.5, italic=True, color="1E40AF")
font_note = Font(name="Calibri", size=9.5, italic=True, color="4B5563")

fill_header_blue = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
fill_result = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid")
fill_input = PatternFill(start_color="F0F9FF", end_color="F0F9FF", fill_type="solid")
fill_gray = PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid")
fill_amber = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")

thin_side = Side(border_style="thin", color="9CA3AF")
thick_bottom = Side(border_style="medium", color="1E40AF")
border_matrix = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

align_center = Alignment(horizontal="center", vertical="center")
align_left = Alignment(horizontal="left", vertical="center")

def style_matrix(ws, start_row, start_col, rows, cols, fill=None, is_bold=False):
    for r in range(start_row, start_row + rows):
        for c in range(start_col, start_col + cols):
            cell = ws.cell(row=r, column=c)
            cell.border = border_matrix
            cell.alignment = align_center
            if fill:
                cell.fill = fill
            if is_bold:
                cell.font = font_bold
            else:
                cell.font = font_regular
            if isinstance(cell.value, float):
                cell.number_format = "0.0000" if abs(cell.value - round(cell.value)) > 1e-4 else "0.0"

# ============================================================
# ЛИСТ 1: ОБУЧАЮЩИЕ ПРИМЕРЫ (I - VI)
# ============================================================
ws1 = wb.create_sheet(title="Примеры I-VI")
ws1.views.sheetView[0].showGridLines = True

ws1.cell(row=1, column=1, value="ЛАБОРАТОРНАЯ РАБОТА: «Матричные операции в MS Excel»").font = font_title
ws1.cell(row=2, column=1, value="Обучающая часть: базовые операции над матрицами с использованием формул массива.").font = font_note

# Пример I: Умножение матриц
ws1.cell(row=4, column=1, value="Пример I: Умножение матриц (функция МУМНОЖ / MMULT)").font = font_section
ws1.cell(row=5, column=1, value="Матрица A (2x3)").font = font_bold
ex1_A = [[1, 2, 3], [4, 5, 6]]
for i in range(2):
    for j in range(3):
        ws1.cell(row=6+i, column=1+j, value=ex1_A[i][j])
style_matrix(ws1, 6, 1, 2, 3, fill=fill_input)

ws1.cell(row=5, column=5, value="Матрица B (3x2)").font = font_bold
ex1_B = [[7, 8], [9, 1], [2, 3]]
for i in range(3):
    for j in range(2):
        ws1.cell(row=6+i, column=5+j, value=ex1_B[i][j])
style_matrix(ws1, 6, 5, 3, 2, fill=fill_input)

ws1.cell(row=10, column=1, value="Результат C = A * B (2x2):").font = font_bold
ex1_C = [[31, 19], [85, 55]]
for i in range(2):
    for j in range(2):
        ws1.cell(row=11+i, column=1+j, value=ex1_C[i][j])
style_matrix(ws1, 11, 1, 2, 2, fill=fill_result, is_bold=True)
ws1.cell(row=13, column=1, value="Формула массива Excel: {=МУМНОЖ(A6:C7; E6:F8)} (Ctrl+Shift+Enter)").font = font_formula

# Пример II: Сложение матриц
ws1.cell(row=15, column=1, value="Пример II: Сложение матриц одинакового размера").font = font_section
ws1.cell(row=16, column=1, value="Матрица X (2x2)").font = font_bold
ex2_X = [[5, 2], [3, 8]]
for i in range(2):
    for j in range(2):
        ws1.cell(row=17+i, column=1+j, value=ex2_X[i][j])
style_matrix(ws1, 17, 1, 2, 2, fill=fill_input)

ws1.cell(row=16, column=5, value="Матрица Y (2x2)").font = font_bold
ex2_Y = [[1, 4], [2, 1]]
for i in range(2):
    for j in range(2):
        ws1.cell(row=17+i, column=5+j, value=ex2_Y[i][j])
style_matrix(ws1, 17, 5, 2, 2, fill=fill_input)

ws1.cell(row=20, column=1, value="Результат X + Y:").font = font_bold
ex2_Sum = [[6, 6], [5, 9]]
for i in range(2):
    for j in range(2):
        ws1.cell(row=21+i, column=1+j, value=ex2_Sum[i][j])
style_matrix(ws1, 21, 1, 2, 2, fill=fill_result, is_bold=True)
ws1.cell(row=23, column=1, value="Формула массива Excel: {=A17:B18 + E17:F18}").font = font_formula

# Пример III: Умножение матрицы на скаляр
ws1.cell(row=25, column=1, value="Пример III: Умножение матрицы на число (скаляр k = 3)").font = font_section
ws1.cell(row=26, column=1, value="Результат 3 * X:").font = font_bold
ex3_Scal = [[15, 6], [9, 24]]
for i in range(2):
    for j in range(2):
        ws1.cell(row=27+i, column=1+j, value=ex3_Scal[i][j])
style_matrix(ws1, 27, 1, 2, 2, fill=fill_result, is_bold=True)
ws1.cell(row=29, column=1, value="Формула: {=3 * A17:B18}").font = font_formula

# Пример IV: Транспонирование
ws1.cell(row=31, column=1, value="Пример IV: Транспонирование матрицы (ТРАНСП / TRANSPOSE)").font = font_section
ws1.cell(row=32, column=1, value="Транспонированная матрица X^T (2x2):").font = font_bold
ex4_T = [[5, 3], [2, 8]]
for i in range(2):
    for j in range(2):
        ws1.cell(row=33+i, column=1+j, value=ex4_T[i][j])
style_matrix(ws1, 33, 1, 2, 2, fill=fill_result, is_bold=True)
ws1.cell(row=35, column=1, value="Формула массива Excel: {=ТРАНСП(A17:B18)}").font = font_formula

# Пример V: Определитель матрицы
ws1.cell(row=37, column=1, value="Пример V: Вычисление определителя (МОПРЕД / MDETERM)").font = font_section
ws1.cell(row=38, column=1, value="det(X) = 5*8 - 2*3 = 34").font = font_bold
ws1.cell(row=39, column=1, value=34.0)
style_matrix(ws1, 39, 1, 1, 1, fill=fill_result, is_bold=True)
ws1.cell(row=40, column=1, value="Формула: =МОПРЕД(A17:B18)").font = font_formula

# Пример VI: Обратная матрица
ws1.cell(row=37, column=5, value="Пример VI: Обратная матрица (МОБР / MINVERSE)").font = font_section
ws1.cell(row=38, column=5, value="Обратная матрица X^(-1) (2x2):").font = font_bold
ex6_Inv = [[8/34, -2/34], [-3/34, 5/34]]
for i in range(2):
    for j in range(2):
        ws1.cell(row=39+i, column=5+j, value=ex6_Inv[i][j])
style_matrix(ws1, 39, 5, 2, 2, fill=fill_result, is_bold=True)
ws1.cell(row=41, column=5, value="Формула массива Excel: {=МОБР(A17:B18)}").font = font_formula

# ============================================================
# ЛИСТ 2: ЗАДАНИЕ ДЛЯ САМОСТОЯТЕЛЬНОЙ РАБОТЫ №1
# ============================================================
ws2 = wb.create_sheet(title="Задание 1")
ws2.views.sheetView[0].showGridLines = True

ws2.cell(row=1, column=1, value="ЗАДАНИЕ ДЛЯ САМОСТОЯТЕЛЬНОЙ РАБОТЫ №1").font = font_title
ws2.cell(row=2, column=1, value="Вычислить выражение:  3*A*B + (A - B)*(A + 2*B^T)").font = font_subtitle

# Исходные матрицы A и B
m1_A = [[2.0, 5.0, -1.0], [0.0, -2.0, 1.0], [1.0, 0.0, 1.0]]
m1_B = [[1.0, -2.0, 0.0], [1.0, 0.0, 2.0], [0.0, 0.0, 3.0]]

ws2.cell(row=4, column=1, value="Матрица A (3x3)").font = font_bold
for i in range(3):
    for j in range(3):
        ws2.cell(row=5+i, column=1+j, value=m1_A[i][j])
style_matrix(ws2, 5, 1, 3, 3, fill=fill_input)

ws2.cell(row=4, column=5, value="Матрица B (3x3)").font = font_bold
for i in range(3):
    for j in range(3):
        ws2.cell(row=5+i, column=5+j, value=m1_B[i][j])
style_matrix(ws2, 5, 5, 3, 3, fill=fill_input)

# Действие 1: A - B
ws2.cell(row=9, column=1, value="Действие 1: Разность матриц (A - B)").font = font_section
m1_A_minus_B = [[1.0, 7.0, -1.0], [-1.0, -2.0, -1.0], [1.0, 0.0, -2.0]]
for i in range(3):
    for j in range(3):
        ws2.cell(row=10+i, column=1+j, value=m1_A_minus_B[i][j])
style_matrix(ws2, 10, 1, 3, 3, fill=fill_gray)
ws2.cell(row=13, column=1, value="Формула: {=A5:C7 - E5:G7}").font = font_formula

# Действие 2: B^T
ws2.cell(row=9, column=5, value="Действие 2: Транспонированная матрица B^T").font = font_section
m1_BT = [[1.0, 1.0, 0.0], [-2.0, 0.0, 0.0], [0.0, 2.0, 3.0]]
for i in range(3):
    for j in range(3):
        ws2.cell(row=10+i, column=5+j, value=m1_BT[i][j])
style_matrix(ws2, 10, 5, 3, 3, fill=fill_gray)
ws2.cell(row=13, column=5, value="Формула: {=ТРАНСП(E5:G7)}").font = font_formula

# Действие 3: A + 2*B^T
ws2.cell(row=15, column=1, value="Действие 3: Матрица (A + 2*B^T)").font = font_section
m1_A_plus_2BT = [[4.0, 7.0, -1.0], [-4.0, -2.0, 1.0], [1.0, 4.0, 7.0]]
for i in range(3):
    for j in range(3):
        ws2.cell(row=16+i, column=1+j, value=m1_A_plus_2BT[i][j])
style_matrix(ws2, 16, 1, 3, 3, fill=fill_gray)
ws2.cell(row=19, column=1, value="Формула: {=A5:C7 + 2*E10:G12}").font = font_formula

# Действие 4: 3*A*B
ws2.cell(row=15, column=5, value="Действие 4: Первое слагаемое  3*A*B").font = font_section
m1_3AB = [[21.0, -12.0, 21.0], [-6.0, 0.0, -3.0], [3.0, -6.0, 9.0]]
for i in range(3):
    for j in range(3):
        ws2.cell(row=16+i, column=5+j, value=m1_3AB[i][j])
style_matrix(ws2, 16, 5, 3, 3, fill=fill_gray)
ws2.cell(row=19, column=5, value="Формула: {=3 * МУМНОЖ(A5:C7; E5:G7)}").font = font_formula

# Действие 5: (A - B)*(A + 2*B^T)
ws2.cell(row=21, column=1, value="Действие 5: Второе слагаемое  (A - B)*(A + 2*B^T)").font = font_section
m1_term2 = [[-25.0, -11.0, -1.0], [3.0, -7.0, -8.0], [2.0, -1.0, -15.0]]
for i in range(3):
    for j in range(3):
        ws2.cell(row=22+i, column=1+j, value=m1_term2[i][j])
style_matrix(ws2, 22, 1, 3, 3, fill=fill_gray)
ws2.cell(row=25, column=1, value="Формула: {=МУМНОЖ(A10:C12; A16:C18)}").font = font_formula

# ИТОГОВЫЙ РЕЗУЛЬТАТ
ws2.cell(row=27, column=1, value="ИТОГОВЫЙ РЕЗУЛЬТАТ: 3*A*B + (A - B)*(A + 2*B^T)").font = font_title
m1_res = [[-4.0, -23.0, 20.0], [-3.0, -7.0, -11.0], [5.0, -7.0, -6.0]]
for i in range(3):
    for j in range(3):
        ws2.cell(row=28+i, column=1+j, value=m1_res[i][j])
style_matrix(ws2, 28, 1, 3, 3, fill=fill_result, is_bold=True)
ws2.cell(row=32, column=1, value="Формула сложения слагаемых: {=E16:G18 + A22:C24}").font = font_formula

# ============================================================
# ЛИСТ 3: ЗАДАНИЕ ДЛЯ САМОСТОЯТЕЛЬНОЙ РАБОТЫ №2
# ============================================================
ws3 = wb.create_sheet(title="Задание 2")
ws3.views.sheetView[0].showGridLines = True

ws3.cell(row=1, column=1, value="ЗАДАНИЕ ДЛЯ САМОСТОЯТЕЛЬНОЙ РАБОТЫ №2").font = font_title
ws3.cell(row=2, column=1, value="Вычислить выражение:  3*B*(A + B) - |A|*(A^(-1) + 2*B^T)").font = font_subtitle

# Исходные матрицы A и B
m2_A = [[2.0, 0.5, 2.0], [1.0, 0.0, -1.0], [1.0, -2.0, 1.0]]
m2_B = [[1.0, -2.0, 1.0], [-0.5, 1.0, 5.0], [-1.0, 2.0, 3.0]]

ws3.cell(row=4, column=1, value="Матрица A").font = font_bold
for i in range(3):
    for j in range(3):
        ws3.cell(row=5+i, column=1+j, value=m2_A[i][j])
style_matrix(ws3, 5, 1, 3, 3, fill=fill_input)

ws3.cell(row=4, column=5, value="Матрица B").font = font_bold
for i in range(3):
    for j in range(3):
        ws3.cell(row=5+i, column=5+j, value=m2_B[i][j])
style_matrix(ws3, 5, 5, 3, 3, fill=fill_input)

# Действие 1: Определитель |A|
ws3.cell(row=9, column=1, value="Действие 1: Определитель |A| (МОПРЕД)").font = font_section
ws3.cell(row=10, column=1, value=-9.0)
style_matrix(ws3, 10, 1, 1, 1, fill=fill_gray, is_bold=True)
ws3.cell(row=11, column=1, value="Формула: =МОПРЕД(A5:C7) -> -9.0").font = font_formula

# Действие 2: Обратная матрица A^(-1)
ws3.cell(row=9, column=5, value="Действие 2: Обратная матрица A^(-1) (МОБР)").font = font_section
m2_invA = [
    [2.0/9.0, 4.5/9.0, 0.5/9.0],
    [2.0/9.0, 0.0/9.0, -4.0/9.0],
    [2.0/9.0, -4.5/9.0, 0.5/9.0]
]
for i in range(3):
    for j in range(3):
        ws3.cell(row=10+i, column=5+j, value=m2_invA[i][j])
style_matrix(ws3, 10, 5, 3, 3, fill=fill_gray)
ws3.cell(row=13, column=5, value="Формула: {=МОБР(A5:C7)}").font = font_formula

# Действие 3: 3*B*(A + B)
ws3.cell(row=15, column=1, value="Действие 3: Первое слагаемое  3*B*(A + B)").font = font_section
m2_term1 = [[6.0, -10.5, -3.0], [-3.0, 5.25, 67.5], [-6.0, 10.5, 51.0]]
for i in range(3):
    for j in range(3):
        ws3.cell(row=16+i, column=1+j, value=m2_term1[i][j])
style_matrix(ws3, 16, 1, 3, 3, fill=fill_gray)
ws3.cell(row=19, column=1, value="Формула: {=3 * МУМНОЖ(E5:G7; (A5:C7 + E5:G7))}").font = font_formula

# Действие 4: |A|*(A^(-1) + 2*B^T)
ws3.cell(row=15, column=5, value="Действие 4: Второе вычитаемое слагаемое  |A|*(A^(-1) + 2*B^T)").font = font_section
m2_term2 = [[-20.0, 4.5, 17.5], [34.0, -18.0, -32.0], [-20.0, -85.5, -54.5]]
for i in range(3):
    for j in range(3):
        ws3.cell(row=16+i, column=5+j, value=m2_term2[i][j])
style_matrix(ws3, 16, 5, 3, 3, fill=fill_gray)
ws3.cell(row=19, column=5, value="Формула: {=A10 * (E10:G12 + 2 * ТРАНСП(E5:G7))}").font = font_formula

# ИТОГОВЫЙ РЕЗУЛЬТАТ
ws3.cell(row=21, column=1, value="ИТОГОВЫЙ РЕЗУЛЬТАТ: 3*B*(A + B) - |A|*(A^(-1) + 2*B^T)").font = font_title
m2_res = [[26.0, -15.0, -20.5], [-37.0, 23.25, 99.5], [14.0, 96.0, 105.5]]
for i in range(3):
    for j in range(3):
        ws3.cell(row=22+i, column=1+j, value=m2_res[i][j])
style_matrix(ws3, 22, 1, 3, 3, fill=fill_result, is_bold=True)
ws3.cell(row=26, column=1, value="Формула: {=A16:C18 - E16:G18} (Разность блоков)").font = font_formula

# ============================================================
# ЛИСТ 4: ЗАДАНИЕ ДЛЯ САМОСТОЯТЕЛЬНОЙ РАБОТЫ №3
# ============================================================
ws4 = wb.create_sheet(title="Задание 3")
ws4.views.sheetView[0].showGridLines = True

ws4.cell(row=1, column=1, value="ЗАДАНИЕ ДЛЯ САМОСТОЯТЕЛЬНОЙ РАБОТЫ №3").font = font_title
ws4.cell(row=2, column=1, value="Вычислить выражение:  3*B^T*(A - B) + |B|*(4*A^T + B^(-1))").font = font_subtitle

# Исходные матрицы A и B
ws4.cell(row=4, column=1, value="Матрица A").font = font_bold
for i in range(3):
    for j in range(3):
        ws4.cell(row=5+i, column=1+j, value=m2_A[i][j])
style_matrix(ws4, 5, 1, 3, 3, fill=fill_input)

ws4.cell(row=4, column=5, value="Матрица B").font = font_bold
for i in range(3):
    for j in range(3):
        ws4.cell(row=5+i, column=5+j, value=m2_B[i][j])
style_matrix(ws4, 5, 5, 3, 3, fill=fill_input)

# Действие 1: Определитель |B|
ws4.cell(row=9, column=1, value="Действие 1: Определитель |B| (МОПРЕД)").font = font_section
ws4.cell(row=10, column=1, value=0.0)
style_matrix(ws4, 10, 1, 1, 1, fill=fill_amber, is_bold=True)
ws4.cell(row=11, column=1, value="Формула: =МОПРЕД(E5:G7) -> 0.0 (Матрица вырождена!)").font = font_formula

# Аналитическое примечание
ws4.cell(row=9, column=4, value="ВАЖНОЕ МАТЕМАТИЧЕСКОЕ ЗАМЕЧАНИЕ О B^(-1):").font = font_bold
ws4.cell(row=10, column=4, value="Так как |B| = 0, обратная матрица B^(-1) не существует (в Excel МОБР возвращает #ЧИСЛО!).").font = font_note
ws4.cell(row=11, column=4, value="Однако при раскрытии скобок: |B|*(4*A^T + B^(-1)) = 4*|B|*A^T + |B|*B^(-1) = 0 + adj(B) = adj(B).").font = font_note
ws4.cell(row=12, column=4, value="Ниже представлены оба варианта решения: аналитический (через присоединенную матрицу) и пошаговый.").font = font_note

# Действие 2: 3*B^T*(A - B)
ws4.cell(row=14, column=1, value="Действие 2: Первое слагаемое  3*B^T*(A - B)").font = font_section
m3_term1 = [[-5.25, 21.0, 18.0], [10.5, -42.0, -36.0], [43.5, -43.5, -105.0]]
for i in range(3):
    for j in range(3):
        ws4.cell(row=15+i, column=1+j, value=m3_term1[i][j])
style_matrix(ws4, 15, 1, 3, 3, fill=fill_gray)
ws4.cell(row=19, column=1, value="Формула: {=3 * МУМНОЖ(ТРАНСП(E5:G7); (A5:C7 - E5:G7))}").font = font_formula

# Действие 3: Присоединенная матрица adj(B)
ws4.cell(row=14, column=5, value="Действие 3: Присоединенная матрица adj(B) = |B|*B^(-1)").font = font_section
m3_adjB = [[-7.0, 8.0, -11.0], [-3.5, 4.0, -5.5], [0.0, 0.0, 0.0]]
for i in range(3):
    for j in range(3):
        ws4.cell(row=15+i, column=5+j, value=m3_adjB[i][j])
style_matrix(ws4, 15, 5, 3, 3, fill=fill_gray)
ws4.cell(row=19, column=5, value="Матрица алгебраических дополнений B* (adjugate matrix)").font = font_formula

# ИТОГОВЫЙ РЕЗУЛЬТАТ (Вариант 1)
ws4.cell(row=21, column=1, value="ИТОГОВЫЙ РЕЗУЛЬТАТ (Вариант 1 - Аналитический через предел/adjugate):").font = font_title
m3_res = [[-12.25, 29.0, 7.0], [7.0, -38.0, -41.5], [43.5, -43.5, -105.0]]
for i in range(3):
    for j in range(3):
        ws4.cell(row=22+i, column=1+j, value=m3_res[i][j])
style_matrix(ws4, 22, 1, 3, 3, fill=fill_result, is_bold=True)
ws4.cell(row=26, column=1, value="Формула: = 3*B^T*(A - B) + adj(B)  (Сумма слагаемых)").font = font_formula

# ИТОГОВЫЙ РЕЗУЛЬТАТ (Вариант 2)
ws4.cell(row=28, column=1, value="ИТОГОВЫЙ РЕЗУЛЬТАТ (Вариант 2 - Формальный пошаговый в Excel):").font = font_section
ws4.cell(row=29, column=1, value="В силу вырожденности матрицы B (|B| = 0), обратная матрица B^(-1) не существует.").font = font_bold
ws4.cell(row=30, column=1, value="В ячейках формулы =МОБР(B) программа Excel возвращает системную ошибку #ЧИСЛО!").font = font_note
ws4.cell(row=31, column=1, value="Поэтому при строгом пошаговом расчете выражение не определено.").font = font_note

# Автоподбор ширины столбцов
for ws in [ws1, ws2, ws3, ws4]:
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val = str(cell.value or "")
            if len(val) > max_len and not val.startswith("ЛАБОРАТОРНАЯ") and not val.startswith("ЗАДАНИЕ"):
                max_len = len(val)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 11)

out1 = os.path.join("lab_excel_matrices", "Лабораторная_работа_Матрицы_Excel.xlsx")
out2 = os.path.join("lab_excel_matrices", "Matrix_Lab_Excel.xlsx")

wb.save(out1)
wb.save(out2)
print("SUCCESS: Workbooks generated successfully.")
